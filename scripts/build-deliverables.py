"""Build submission artifacts from Markdown and actual captured evidence.
Run from any directory with the packages in requirements-deliverables.txt.
Does not read .env, keys, databases or browser sessions.
"""
from pathlib import Path
import json,re,math,html,textwrap,zipfile,hashlib
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor as DocColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph,Table,TableStyle,Image as RLImage,Spacer,PageBreak,SimpleDocTemplate
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pptx import Presentation
from pptx.util import Inches as PIn,Pt as PPt
from pptx.dml.color import RGBColor
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'deliverables'; DIAG=OUT/'diagrams'
OUT.mkdir(exist_ok=True);DIAG.mkdir(exist_ok=True)
INK='#193e37';GREEN='#24594a';MUTED='#62756c';BG='#f7f8f3';LINE='#cbd6c8';GOLD='#bd975d'
font_candidates=[Path('/System/Library/Fonts/Supplemental/Arial.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
FONT=next((p for p in font_candidates if p.exists()),None)
BOLD=FONT.with_name('Arial Bold.ttf') if FONT and FONT.name=='Arial.ttf' else (FONT.with_name('DejaVuSans-Bold.ttf') if FONT else None)
if FONT: pdfmetrics.registerFont(TTFont('Body',str(FONT)));pdfmetrics.registerFont(TTFont('BodyBold',str(BOLD if BOLD and BOLD.exists() else FONT)))
BODY='Body' if FONT else 'Helvetica';BOLDPDF='BodyBold' if FONT else 'Helvetica-Bold'
def pf(size,bold=False):return ImageFont.truetype(str(BOLD if bold and BOLD and BOLD.exists() else FONT),size) if FONT else ImageFont.load_default(size=size)
class Diagram:
 def __init__(self,title,subtitle,h=700):
  self.h=h; self.im=Image.new('RGB',(1600,h),BG);self.d=ImageDraw.Draw(self.im);self.svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="{h}" viewBox="0 0 1600 {h}"><rect width="1600" height="{h}" fill="{BG}"/>']
  self.text(50,30,title,36,True);self.text(50,82,subtitle,22)
 def text(self,x,y,text,size=24,bold=False,color=INK):
  for i,line in enumerate(str(text).split('\n')):
   yy=y+i*(size+9);self.d.text((x,yy),line,font=pf(size,bold),fill=color)
   self.svg.append(f'<text x="{x}" y="{yy+size}" font-family="Arial,sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{color}">{html.escape(line)}</text>')
 def box(self,x,y,w,h,title,body='',fill='white'):
  self.d.rounded_rectangle((x,y,x+w,y+h),radius=16,fill=fill,outline=LINE,width=2)
  self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="16" fill="{fill}" stroke="{LINE}" stroke-width="2"/>')
  self.text(x+20,y+18,title,25,True)
  if body:self.text(x+20,y+62,body,21)
 def arrow(self,x1,y1,x2,y2,label='',lx=None,ly=None):
  self.d.line((x1,y1,x2,y2),fill=GREEN,width=4)
  a=math.atan2(y2-y1,x2-x1);pts=[(x2,y2),(x2-16*math.cos(a-.45),y2-16*math.sin(a-.45)),(x2-16*math.cos(a+.45),y2-16*math.sin(a+.45))]
  self.d.polygon(pts,fill=GREEN);self.svg.append(f'<path d="M{x1},{y1} L{x2},{y2}" fill="none" stroke="{GREEN}" stroke-width="4"/><polygon points="'+ ' '.join(f'{x},{y}' for x,y in pts)+f'" fill="{GREEN}"/>')
  if label:self.text(lx if lx is not None else (x1+x2)/2-40,ly if ly is not None else (y1+y2)/2-34,label,20)
 def save(self,name,mermaid):
  self.im.save(DIAG/f'{name}.png');(DIAG/f'{name}.svg').write_text(''.join(self.svg)+'</svg>');(DIAG/f'{name}.mmd').write_text(mermaid+'\n')

d=Diagram('01  System architecture','Implemented prototype · custody stays on-chain; private identities stay in SQLite')
d.box(50,150,340,150,'Browser / Next.js','Dashboard + account forms\nSame-origin API requests')
d.box(630,150,350,150,'Python / FastAPI','Sessions + wallet proof\nReceipt verification')
d.box(1210,150,340,150,'SQLite database','Accounts + sessions\nChallenges + event journal')
d.box(50,470,340,150,'MetaMask','User-controlled keys\nMessages + transactions')
d.box(630,470,350,150,'Sepolia RPC','Independent chain reads\nTransaction broadcast')
d.box(1210,470,340,150,'Savings.sol','Principal + interest\nOwner-funded reserve')
d.arrow(390,225,630,225,'/api',465,184);d.arrow(980,225,1210,225,'persist',1040,184)
d.arrow(210,300,210,470,'sign',225,373);d.arrow(390,545,630,545,'signed tx',427,505)
d.arrow(805,300,805,470,'read receipts',823,373);d.arrow(980,545,1210,545,'read / execute',1000,505)
d.save('01-architecture','''flowchart LR
 B[Browser / Next.js] -->|same origin /api| API[Python FastAPI]
 API --> DB[(SQLite)]
 B -->|wallet prompts| M[MetaMask]
 M -->|signed transaction| RPC[Sepolia RPC]
 API -->|read balances and verify receipts| RPC
 RPC --> V[Savings.sol]''')
d=Diagram('02  Smart contract interaction','Agreed focused scope: one vault · four genuine transaction types')
d.box(50,180,330,160,'Customer wallet','deposit(value)\nwithdraw(amount) / claim()')
d.box(630,160,470,355,'Savings.sol','accounts[wallet]\ntotalPrincipal / rateBps / owner\ncheckpoint before balance change\nreserve = balance - principal\nclaim only from surplus',fill='#e8eee0')
d.box(50,450,330,150,'Owner wallet','fund(value)\nNo member withdrawal right')
d.box(1250,180,300,150,'Event logs','Deposited / Withdrawn\nInterestPaid / Funded')
d.box(1250,450,300,150,'Python API','Read balances\nVerify event receipts')
d.arrow(380,255,630,255,'signed tx',425,213);d.arrow(380,525,630,450,'fund',445,452)
d.arrow(1100,240,1250,240,'emit',1140,205);d.arrow(1400,330,1400,450,'read',1420,375)
d.text(650,555,'ETH payouts return directly to caller wallet.',22)
d.save('02-contract-interaction','''flowchart LR
 C[Customer wallet] -->|deposit / withdraw / claim| S[Savings: only implemented contract]
 O[Immutable owner wallet] -->|fund| S
 S -->|ETH withdrawal or interest| C
 S --> E[Four event types]
 A[API] -->|read state and verify receipts| S
 E --> A''')
d=Diagram('03  Persistent data model','Implemented SQLite tables · PK = primary key · FK = foreign key · UQ = unique')
d.box(50,150,425,205,'audit_log','id PK / user_id FK\naction / detail / created')
d.arrow(600,260,475,260,'1 : many',495,213)
d.box(600,150,400,220,'users','id PK / email UQ\npassword: salted scrypt\nrole: user / admin / auditor\nwallet UQ')
d.box(50,430,425,205,'sessions','token PK: SHA-256 hash\nuser_id FK -> users.id\nexpires: milliseconds')
d.box(585,430,430,205,'challenges','user_id PK / FK -> users.id\nmessage: nonce + account\nexpires: milliseconds')
d.box(1120,430,430,220,'transactions','(hash, log_index) composite PK\nuser_id FK / kind / amount\nblock / block_hash / block_time\ncreated: recording time')
d.arrow(600,370,260,430,'1 : many',360,380);d.arrow(800,370,800,430,'1 : 0..1',825,385);d.arrow(1000,260,1335,430,'1 : many',1140,335)
d.save('03-data-model','''erDiagram
 users ||--o{ sessions : owns
 users ||--o| challenges : has
 users ||--o{ transactions : records
 users ||--o{ audit_log : records
 users { INTEGER id PK
 TEXT email UK
 TEXT password
 TEXT role
 TEXT wallet UK }
 sessions { TEXT token PK
 INTEGER user_id FK
 INTEGER expires }
 challenges { INTEGER user_id PK,FK
 TEXT message
 INTEGER expires }
 transactions { TEXT hash PK
 INTEGER log_index PK
 INTEGER user_id FK
 TEXT kind
 TEXT amount
 INTEGER block
 TEXT created }
 audit_log { INTEGER id PK
 INTEGER user_id FK
 TEXT action
 TEXT detail
 TEXT created }''')
d=Diagram('04  Wallet and deposit process','Message proof and an ETH transaction are separate approvals',h=540)
for x,title,body in [(50,'1  Sign in','Password session'),(570,'2  Link wallet','Nonce + signed proof'),(1090,'3  Submit deposit','Sepolia + amount + fee')]: d.box(x,150,460,115,title,body)
d.arrow(510,208,570,208);d.arrow(1030,208,1090,208);d.arrow(1320,265,1320,345)
for x,title,body in [(1090,'4  Confirm on chain','Wait two confirmations'),(570,'5  Verify receipt','API checks event + wallet'),(50,'6  Store and refresh','Journal + live balances')]:d.box(x,345,460,115,title,body)
d.arrow(1090,402,1030,402);d.arrow(570,402,510,402)
d.text(50,492,'On failure: retain hash, inspect explorer, retry import after confirmation. No fabricated success.',20)
d.save('04-process-flow','''flowchart LR
 A[Password sign in] --> B[Issue wallet challenge]
 B --> C[Sign and verify wallet message]
 C --> D[Review MetaMask Sepolia deposit]
 D --> E[Wait for two confirmations]
 E --> F[API verifies receipt and wallet event]
 F --> G[Idempotent journal insert]
 G --> H[Refresh balances]
 E -->|browser interrupted| I[Import retained hash]
 I --> F''')
d=Diagram('05  Dashboard wireframe','Current layout · three roles, recorded-activity chart and fee preview',h=640)
d.box(50,135,220,460,'Sidebar','Brand + network\nOverview\nTransactions\nAbout')
d.box(300,135,1250,60,'Heading / account / connected wallet')
for x,title in [(300,'Principal'),(725,'Unclaimed interest'),(1150,'Fixed APR')]:d.box(x,215,400,70,title)
d.box(300,305,1250,85,'14-day activity chart + accessible data table','Deposits / withdrawals / interest / funding · recording dates')
d.box(300,410,610,80,'Deposit / withdraw / claim','Fee cap preview, then wallet signature')
d.box(940,410,610,80,'Reserve + role controls','Owner funding; auditors inspect only')
d.box(300,515,1250,80,'History, filters, pagination, CSV and hash import','Privileged users can inspect the account audit trail')
d.save('05-ui-wireframe','''flowchart TB
 A[Heading and account wallet] --> B[Principal / interest / APR]
 B --> C[14-day recorded activity chart and data table]
 C --> D[Transaction amount and fee preview]
 C --> E[Reserve and role controls]
 D --> F[History: filter, paginate, CSV and hash import]
 E --> F
 F --> G[Admin and auditor account audit trail]''')

# Create testing analysis from raw evidence, without assuming success.
checks=json.loads((OUT/'evidence/checks.json').read_text());gas=json.loads((OUT/'evidence/gas-benchmark.json').read_text());off,on=gas['results']
rows='\n'.join(f"| {c['check']} | {'PASS' if c['exitCode']==0 else 'FAIL'} | {c['exitCode']} | evidence/{c['log']} |" for c in checks['checks'])
gasrows='\n'.join(f"| {a['name']} | {a['estimatedGas']} | {a['actualGas']} | {b['estimatedGas']} | {b['actualGas']} | {(1-int(b['actualGas'])/int(a['actualGas']))*100:.2f}% |" for a,b in zip(off['operations'],on['operations']))
testing=f'''# Testing, Gas and Evaluation Report
Recorded 18 October 2026 · Zeran Yao (G2509029A)

## 1. Reproducibility and scope
Run `python3 scripts/collect-evidence.py` from the project root. Raw stdout/stderr, exit codes and UTC start times are stored in evidence/. Host: {checks['host']}. The backend uses Python 3.12/FastAPI and pytest; frontend/contract tooling uses Node 24. The Python suite has 16 cases and the EVM suite has three tests; assertions are not counted as independent cases. This is functional evidence, not a security certification.

| Check | Result | Exit code | Raw evidence |
| --- | --- | --- | --- |
{rows}

The Python runtime Docker image and combined-container Origin/session smoke are checked separately by scripts/docker-smoke.py. The browser workflow is checked separately by scripts/browser-smoke.py. Their raw evidence files disclose the exact local environment. No public Render or Sepolia result is claimed.

## 2. Tested scenarios and limitations
Backend cases cover actual Node-generated password compatibility; additive legacy migration; sessions and expiry; invitation escalation; wallet proof, expiry and replay; immutable/unique wallets; populated journal isolation; auditor read-only access and email redaction; pagination/filtering; CSV formula safety; precise wei analytics; decoding and rejecting untrusted receipts; Origin checks and request limits. Tests use temporary databases, real signatures and controlled RPC data.

Three EVM tests check funded accounting and invalid amounts; principal withdrawal with an empty reserve and retained interest; and top-up/partial-withdrawal accounting without retroactive interest. They use Ganache and simulated time. Malicious callbacks, failed recipients, competing claim fairness and arbitrary-sequence invariants are not covered.

Build/lint validate compilation, types and code rules. The separate Chromium/local-EVM workflow checks a rendered dashboard, three transaction previews/executions, confirmed journal imports, filters, chart data table, CSV and auditor controls, including a 390px layout. Its wallet is a test EIP-1193 adapter, not the MetaMask extension. No stakeholder feedback or actual user timing was gathered.

## 3. Controlled gas experiment
Compiler 0.8.30, EVM Shanghai, optimizer comparison off versus on (200 runs), fixed initial simulated chain time, fresh deployment per configuration, 500-basis-point rate. Sequence: deposit 1 ETH; owner funds 0.2 ETH; advance one year; claim; withdraw 1 ETH. All measurements are gas units, not ETH or milliseconds.

| Method | Estimate off | Actual off | Estimate on | Actual on | Actual reduction |
| --- | --- | --- | --- | --- | --- |
{gasrows}

Creation bytecode: {off['creationBytecodeBytes']} bytes without optimizer and {on['creationBytecodeBytes']} with it ({(1-on['creationBytecodeBytes']/off['creationBytecodeBytes'])*100:.2f}% smaller). Deployment gas: {off['deploymentGas']} versus {on['deploymentGas']} ({(1-int(on['deploymentGas'])/int(off['deploymentGas']))*100:.2f}% lower). Source: evidence/gas-benchmark.json. One run per configuration supports only a descriptive comparison.

Receipt gas can be below estimate because state changes, storage refunds and estimation strategy affect the result. Results apply to this particular state sequence. Repeated deposits and other network forks may differ. Public-network latency, current gas price and fiat cost are not measured. At a hypothetical 10 gwei gas price, an optimized claim of {on['operations'][2]['actualGas']} gas would cost {int(on['operations'][2]['actualGas'])*10/1e9:.8f} ETH. This assumption is illustrative, not a live quote. Ethereum explains the gas/price relationship in https://ethereum.org/developers/docs/gas/.

## 4. Measurable impact and validation
Observed: selected local accounting and authentication assertions pass, principal becomes zero after full withdrawal, and optimizer-enabled compilation reduces gas in this experiment. Not observed: improved financial access, lower user cost than existing services, higher trust, user satisfaction, public settlement time or scalability under concurrent load. Contextual research must not be reported as a result of this prototype.

Proposed pilot: recruit five consenting wallet-literate participants, ask each to link a wallet, deposit test ETH, identify reserve risk and locate a receipt, then record unaided completion, elapsed time, errors and feedback. Target: at least four complete unaided and can explain why a claim might fail. This pilot is a plan; no participant data or results are supplied.

## 5. Pending tests and release decision
Under the owner-authorised relaxed scope, one contract and four workflows are intentional. Auditor access, charting and fee previews are implemented. Next add hostile-recipient and disk restore tests, later-reorganisation reconciliation and load tests. Run real MetaMask/public-host checks and capture actual public receipts. Dependency advisories in development tools were reported during installation; no independent audit or vulnerability remediation assessment is claimed.

Release decision: local prototype checks support the agreed focused scope; original numeric thresholds are disclosed and enterprise production readiness is not established. See COMPLIANCE_MATRIX.md and DEPLOYMENT_RECORD.md for missing evidence.
'''
(OUT/'TESTING_REPORT.md').write_text(testing)
# Include the key measurements in the bounded technical report.
r=OUT/'TECHNICAL_REPORT.md';s=r.read_text();start='The frontend now estimates gas through the connected wallet provider,'
summary=f"Measured deployment gas fell from {off['deploymentGas']} to {on['deploymentGas']} ({(1-int(on['deploymentGas'])/int(off['deploymentGas']))*100:.2f}%). Optimized actual gas was "+', '.join(f"{a['name']}: {a['actualGas']}" for a in on['operations'])+'. '
s=re.sub(r'Measured deployment gas fell from [^\n]*\n\n','',s)
s=s.replace(start,summary+'\n\n'+start);r.write_text(s)

# Common Markdown block reader: headings, paragraphs, lists, tables and diagrams.
def blocks(text):
 lines=text.splitlines();i=0
 while i<len(lines):
  l=lines[i].strip()
  if not l:i+=1;continue
  if l=='<!-- PAGE -->':yield ('page',None);i+=1;continue
  if l.startswith('!['):
   m=re.match(r'!\[(.*?)\]\((.*?)\)',l);yield ('image',m.groups());i+=1;continue
  if l.startswith('|'):
   table=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r':?-+:?',c) for c in cells):table.append(cells)
    i+=1
   yield('table',table);continue
  if l.startswith('#'):
   m=re.match(r'(#+)\s+(.*)',l);yield('h'+str(len(m[1])),m[2]);i+=1;continue
  if re.match(r'^(\d+\.|-) ',l):yield('p',l);i+=1;continue
  parts=[l];i+=1
  while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','![','<!--')):
   if re.match(r'^(\d+\.|-) ',lines[i]):break
   parts.append(lines[i].strip());i+=1
  yield('p',' '.join(parts))
def clean(t):return t.replace('`','').replace('**','')
def esc(t):return html.escape(clean(t))
P=ParagraphStyle('body',fontName=BODY,fontSize=10,leading=14.4,textColor=colors.HexColor(INK),spaceAfter=9,splitLongWords=True)
SM=ParagraphStyle('small',parent=P,fontSize=8,leading=10.5,spaceAfter=0)
styles={'p':P,'h1':ParagraphStyle('h1',parent=P,fontName=BOLDPDF,fontSize=22,leading=26,spaceAfter=16),'h2':ParagraphStyle('h2',parent=P,fontName=BOLDPDF,fontSize=14,leading=18,spaceAfter=12),'h3':ParagraphStyle('h3',parent=P,fontName=BOLDPDF,fontSize=11,leading=15,spaceBefore=6,spaceAfter=8)}
def page_decoration(c,doc):
 w,h=A4;c.setStrokeColor(colors.HexColor(LINE));c.line(44,h-42,w-44,h-42);c.setFont(BODY,8);c.setFillColor(colors.HexColor(MUTED));c.drawString(44,h-31,'SAVING  /  ENTERPRISE FINANCIAL DAPP');c.drawRightString(w-44,h-31,'Zeran Yao · G2509029A');c.drawString(44,27,'Prototype evidence · 18 October 2026');c.drawRightString(w-44,27,str(doc.page))
def build_pdf(stem):
 story=[]
 for kind,val in blocks((OUT/(stem+'.md')).read_text()):
  if kind=='page':story.append(PageBreak())
  elif kind=='image':
   caption,path=val;im=Image.open(OUT/path);w=480;h=w*im.height/im.width
   if stem=='TECHNICAL_REPORT' and h>140:w=w*140/h;h=140
   story.extend([RLImage(str(OUT/path),width=w,height=h),Paragraph(esc(caption),SM),Spacer(1,9)])
  elif kind=='table':
   n=len(val[0]);t=Table([[Paragraph(esc(c),SM) for c in row] for row in val],colWidths=[480/n]*n,hAlign='LEFT',repeatRows=1)
   t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e6edde')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor(LINE))]));story.extend([t,Spacer(1,12)])
  else:story.append(Paragraph(esc(val),styles.get(kind,P)))
 doc=SimpleDocTemplate(str(OUT/(stem+'.pdf')),pagesize=A4,rightMargin=55,leftMargin=55,topMargin=60,bottomMargin=48,title=stem.replace('_',' ').title(),author='Zeran Yao (G2509029A)')
 doc.build(story,onFirstPage=page_decoration,onLaterPages=page_decoration)
def build_docx(stem):
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
 normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.1
 for nm,sz in [('Title',24),('Heading 1',20),('Heading 2',14),('Heading 3',11)]:doc.styles[nm].font.name='Arial';doc.styles[nm].font.size=Pt(sz);doc.styles[nm].font.color.rgb=DocColor.from_string(INK[1:])
 sec.header.paragraphs[0].text='SAVING  /  Zeran Yao · G2509029A'
 foot=sec.footer.paragraphs[0];foot.text='Prototype evidence · 18 October 2026  |  '
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
 for kind,val in blocks((OUT/(stem+'.md')).read_text()):
  if kind=='page':doc.add_page_break()
  elif kind=='image':
   image=Image.open(OUT/val[1]);w=6.65;h=w*image.height/image.width
   if stem=='TECHNICAL_REPORT' and h>1.94:w=w*1.94/h
   doc.add_picture(str(OUT/val[1]),width=Inches(w));p=doc.add_paragraph(val[0]);p.runs[0].italic=True;p.runs[0].font.size=Pt(8)
  elif kind=='table':
   t=doc.add_table(rows=0,cols=len(val[0]));t.style='Light Shading Accent 1'
   for row in val:
    cells=t.add_row().cells
    for c,v in zip(cells,row):c.text=clean(v)
   for row in t.rows:
    for cell in row.cells:
     for p in cell.paragraphs:
      p.paragraph_format.space_after=Pt(3)
      for run in p.runs:run.font.size=Pt(8)
   doc.add_paragraph().paragraph_format.space_after=Pt(2)
  elif kind.startswith('h'):doc.add_heading(clean(val),min(int(kind[1:]),3))
  else:doc.add_paragraph(clean(val))
 doc.core_properties.title=stem.replace('_',' ').title();doc.core_properties.author='Zeran Yao';doc.save(OUT/(stem+'.docx'))

for stem in ['TECHNICAL_REPORT','USER_MANUAL','INDIVIDUAL_CONTRIBUTION','TESTING_REPORT','DEPLOYMENT_RECORD','DEMO_SCRIPT','COMPLIANCE_MATRIX']:
 build_pdf(stem)
 if stem in ['TECHNICAL_REPORT','USER_MANUAL','INDIVIDUAL_CONTRIBUTION']:build_docx(stem)

slides=[
 ('Saving','Transparent, wallet-controlled savings', ['Enterprise Financial DApp · Sepolia prototype','Zeran Yao · G2509029A · Sole contributor','Python API · three roles · charts · fee preview · testnet only.'],None,'Introduce the student identity and sponsor-funded savings scope. Do not describe this as a bank or a production investment product.'),
 ('01 / Problem','Can members verify the rules?', ['Proposed setting: a crypto-literate membership savings incentive','Member needs: visible principal, executable withdrawals and explainable rewards','Sponsor funds interest; there is no external yield source','Problem demand has not been validated with a real organisation.'],None,'Explain why an operator spreadsheet cannot independently enforce custody rules. Label the use case as proposed rather than interviewed.'),
 ('02 / Evidence','Context supports investigation, not adoption claims', ['World Bank Findex 2025: formal saving reached 40% of adults in developing economies in 2024 [1]','This does not establish demand for a crypto savings app.','BIS Papers 156 discusses DeFi asymmetry and stability risks [2].','A bank or ordinary portal may be the better option for many users.'],None,'Cite report references 1 and 2. Explain the limited applicability of global inclusion statistics to an existing-wallet target group.'),
 ('03 / Architecture','Identity and custody are separate','', '01-architecture','Walk through the browser, MetaMask, API, SQLite and vault. A password session controls portal access; wallet signatures control funds.'),
 ('04 / Accounting','Four business transaction types', ['Deposit → increase principal','Withdraw → reduce principal and send ETH','Claim → pay accrued interest from reserve','Fund → owner supplies reserve; no admin access to member principal'], '02-contract-interaction','Show the single implemented contract. At five percent simple interest, one ETH over a 365-day year earns 0.05 ETH. Reserve exhaustion delays interest claims.'),
 ('05 / Live workflow','Sign, confirm, verify, record','', '04-process-flow','Run the customer demonstration. Inspect the MetaMask network, destination and gas fee. If the transaction is pending, preserve its hash and use a clearly labelled earlier receipt.'),
 ('06 / Test evidence','Local checks pass; scope is bounded', ['Python: 16 cases, including migration, auditors, receipts and scoped history','EVM: three tests covering reserve, accrual, top-ups and withdrawals','Solidity compile, frontend lint and production build pass','No live MetaMask, stakeholder pilot or independent security audit claimed.'],None,'Use captured logs in evidence/. Explain that assertion counts are not independent test counts. Simulated time is not real Sepolia elapsed time.'),
 ('07 / Gas','Optimizer comparison, measured locally', [f"Deployment gas: {off['deploymentGas']} → {on['deploymentGas']} ({(1-int(on['deploymentGas'])/int(off['deploymentGas']))*100:.2f}% reduction)",f"Creation bytecode: {off['creationBytecodeBytes']} → {on['creationBytecodeBytes']} bytes",'Optimized receipt gas: '+', '.join(f"{a['name']} {a['actualGas']}" for a in on['operations']),'One controlled Ganache scenario; no current gas-price or latency claim'],None,'Explain optimizer off/on, 200 runs, compiler 0.8.30, Shanghai target and identical operation sequence. Fees depend on gas price. See testing report.'),
 ('08 / Security','Principal protection is only one control', ['Interest claims cannot consume principal','Reentrancy guard and accounting before ETH transfer','Salted password hashing, expiring cookies and one-use wallet challenges','Open risks: finite reserve, lost wallet keys, RPC trust and missing recovery'],None,'Discuss public ledger privacy and email-to-wallet linkage. Passing tests do not establish audited safety or sustainable funding.'),
 ('09 / Deployment','Docker-ready; public evidence is pending', ['Next.js standalone + Python/FastAPI + SQLite persistent disk','Python runtime Docker build and browser-origin/session smoke passed','No public Render URL, Sepolia address or GitHub history supplied','Live deployment, restart persistence and performed demonstration still required.'],None,'Show the actual deployment register. Do not substitute localhost for the required public URL. Show real explorer evidence only if available.'),
 ('10 / Scope','Focused scope, transparent limits', ['Implemented: Customer, Administrator and read-only Auditor','Charting, paginated history, CSV, audit trail and fee preview','One vault and four meaningful workflows retained under owner-authorised flexibility','Public deployment, real MetaMask acceptance and performed demo remain pending'],None,'Explain the agreed flexible scope and practical improvements. Retain the original-count comparison without adding unrelated functionality.'),
 ('11 / Contribution','An inspectable prototype, with a clear next step', ['Zeran Yao · G2509029A · 100% sole project contribution','Artifacts: source, 14-page report, five diagrams, tests, manual and slides','Next: finish mandatory workflows, deploy, validate with users and record live evidence','No return, adoption or enterprise-readiness claim is made.'],None,'Conclude the ten-minute demonstration. Answer questions from actual evidence, identify limitations and do not invent contribution history.')
]
prs=Presentation();prs.slide_width=PIn(13.333);prs.slide_height=PIn(7.5)
pc=canvas.Canvas(str(OUT/'PRESENTATION.pdf'),pagesize=(960,540));pc.setTitle('Saving — Presentation');pc.setAuthor('Zeran Yao')
def ptext(slide,x,y,w,h,text,size,color=INK,bold=False):
 shape=slide.shapes.add_textbox(PIn(x),PIn(y),PIn(w),PIn(h));tf=shape.text_frame;tf.word_wrap=True
 for i,line in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=PPt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color[1:])
 return shape
for idx,(label,title,bullets,diagram,notes) in enumerate(slides,1):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(BG[1:])
 ptext(slide,.55,.35,12,.5,label.upper(),13,GREEN,True);ptext(slide,.55,1.05,12,1,title,30,INK,True)
 pc.setFillColor(colors.HexColor(BG));pc.rect(0,0,960,540,fill=1,stroke=0);pc.setFillColor(colors.HexColor(GREEN));pc.setFont(BOLDPDF,12);pc.drawString(40,500,label.upper());pc.setFont(BOLDPDF,27);pc.setFillColor(colors.HexColor(INK));pc.drawString(40,443,title)
 if diagram and not bullets:
  im=Image.open(DIAG/(diagram+'.png'));w=12;h=w*im.height/im.width
  slide.shapes.add_picture(str(DIAG/(diagram+'.png')),PIn(.6),PIn(2.05),width=PIn(w),height=PIn(h))
  pc.drawImage(str(DIAG/(diagram+'.png')),40,75,width=880,height=330,preserveAspectRatio=True,anchor='c')
 else:
  for j,b in enumerate(bullets):
   ptext(slide,.7,2.15+j*.9,11.9,.85,'• '+b,19)
   par=Paragraph(esc('• '+b),ParagraphStyle('slide',fontName=BODY,fontSize=18,leading=23,textColor=colors.HexColor(INK)));pw,ph=par.wrap(870,100);par.drawOn(pc,50,365-j*66-ph+23)
 ptext(slide,.6,7,11,.3,'Zeran Yao · G2509029A  /  Sepolia prototype',10,MUTED);ptext(slide,12.2,7,.5,.3,str(idx),10,MUTED)
 slide.notes_slide.notes_text_frame.text=notes
 pc.setFont(BODY,9);pc.setFillColor(colors.HexColor(MUTED));pc.drawString(40,22,'Zeran Yao · G2509029A  /  Sepolia prototype');pc.drawRightString(920,22,str(idx));pc.showPage()
prs.save(OUT/'PRESENTATION.pptx');pc.save()
(OUT/'PRESENTATION_NOTES.md').write_text('# Presentation speaker notes\n\n'+'\n\n'.join(f'## Slide {i}: {x[1]}\n{x[4]}' for i,x in enumerate(slides,1)))
report_pdf=PdfReader(OUT/'TECHNICAL_REPORT.pdf');assert len(report_pdf.pages)==14,f'Report must be 14 pages, got {len(report_pdf.pages)}'
assert len(PdfReader(OUT/'PRESENTATION.pdf').pages)==12
assert len(Presentation(OUT/'PRESENTATION.pptx').slides)==12
for i,p in enumerate(report_pdf.pages,1):assert len(p.extract_text())>150,f'Empty report page {i}'
print('Generated report (14 PDF pages), 12-slide deck, five diagrams and supporting documents.')
